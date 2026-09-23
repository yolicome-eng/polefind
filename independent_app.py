"""Windows desktop app: no browser, no ElecMap requests."""
import os,queue,re,threading,tkinter as tk
from tkinter import filedialog,messagebox,ttk
from openpyxl import load_workbook
from independent_core import REGION_EPSG,KNOWN_ORIGINS,CODE,calibrate,coordinates,reverse_address

class App:
    def __init__(self):
        self.root=tk.Tk();self.root.title('전산화번호 주소 찾기 — 독립형');self.root.geometry('690x465')
        self.root.resizable(False,False);self.events=queue.Queue();self.origins=dict(KNOWN_ORIGINS)
        self.region=tk.StringVar(value='경기/충청');self.key=tk.StringVar();self.code=tk.StringVar()
        self.sample_path=tk.StringVar();self.input_path=tk.StringVar();self.result=tk.StringVar(value='전산화번호를 입력해 주세요.')
        self.status=tk.StringVar(value='ElecMap을 열지 않습니다. 인터넷 주소 조회에 카카오 REST API 키가 필요합니다.')
        frame=ttk.Frame(self.root,padding=18);frame.pack(fill='both',expand=True)
        ttk.Label(frame,text='전산화번호 → 주소',font=('맑은 고딕',18,'bold')).grid(row=0,column=0,columnspan=3,sticky='w',pady=(0,14))
        ttk.Label(frame,text='권역').grid(row=1,column=0,sticky='w')
        ttk.Combobox(frame,textvariable=self.region,values=list(REGION_EPSG),state='readonly',width=20).grid(row=1,column=1,sticky='w',pady=5)
        ttk.Label(frame,text='카카오 REST API 키').grid(row=2,column=0,sticky='w')
        ttk.Entry(frame,textvariable=self.key,show='•',width=48).grid(row=2,column=1,columnspan=2,sticky='w',pady=5)
        ttk.Label(frame,text='키는 저장하지 않으며 주소 API에만 전송됩니다.',foreground='#555').grid(row=3,column=1,columnspan=2,sticky='w')
        ttk.Label(frame,text='전산화번호').grid(row=4,column=0,sticky='w',pady=(20,5))
        ttk.Entry(frame,textvariable=self.code,width=25,font=('맑은 고딕',12)).grid(row=4,column=1,sticky='w',pady=(20,5))
        self.find_btn=ttk.Button(frame,text='주소 찾기',command=self.find);self.find_btn.grid(row=4,column=2,sticky='w',pady=(20,5))
        ttk.Label(frame,textvariable=self.result,wraplength=625,font=('맑은 고딕',11)).grid(row=5,column=0,columnspan=3,sticky='w',pady=10)
        ttk.Separator(frame).grid(row=6,column=0,columnspan=3,sticky='ew',pady=14)
        ttk.Label(frame,text='샘플 엑셀').grid(row=7,column=0,sticky='w')
        ttk.Entry(frame,textvariable=self.sample_path,width=49).grid(row=7,column=1,sticky='w')
        ttk.Button(frame,text='선택',command=lambda:self.pick(self.sample_path)).grid(row=7,column=2,sticky='w')
        ttk.Button(frame,text='다른 권역 보정·검증',command=self.fit).grid(row=8,column=1,sticky='w',pady=7)
        ttk.Label(frame,text='입력 엑셀').grid(row=9,column=0,sticky='w')
        ttk.Entry(frame,textvariable=self.input_path,width=49).grid(row=9,column=1,sticky='w')
        ttk.Button(frame,text='선택',command=lambda:self.pick(self.input_path)).grid(row=9,column=2,sticky='w')
        self.batch_btn=ttk.Button(frame,text='엑셀 주소 입력',command=self.batch);self.batch_btn.grid(row=10,column=1,sticky='w',pady=8)
        ttk.Label(frame,textvariable=self.status,wraplength=650,foreground='#333').grid(row=11,column=0,columnspan=3,sticky='w',pady=(12,0))
        frame.columnconfigure(1,weight=1);self.root.after(100,self.poll)
    def pick(self,target):
        p=filedialog.askopenfilename(filetypes=[('Excel','*.xlsx')]);
        if p:target.set(p)
    def poll(self):
        while not self.events.empty():
            kind,value=self.events.get()
            if kind=='status':self.status.set(value)
            elif kind=='result':self.result.set(value)
            elif kind=='error':self.status.set('오류: '+value);messagebox.showerror('주소 조회 오류',value)
            elif kind=='done':self.status.set(value);messagebox.showinfo('완료',value)
            elif kind=='origin':self.origins[self.region.get()]=value
            elif kind=='enable':self.find_btn.config(state='normal');self.batch_btn.config(state='normal')
        self.root.after(100,self.poll)
    def run(self,fn):
        if not self.key.get().strip():messagebox.showwarning('API 키','카카오 REST API 키를 입력해 주세요.');return
        self.find_btn.config(state='disabled');self.batch_btn.config(state='disabled')
        key=self.key.get().strip();region=self.region.get()
        def task():
            try:fn(region,key)
            except Exception as e:self.events.put(('error',str(e)))
            finally:self.events.put(('enable',None))
        threading.Thread(target=task,daemon=True).start()
    def sample_rows(self,region):
        p=self.sample_path.get()
        if not os.path.isfile(p):raise ValueError('샘플 엑셀 파일을 선택해 주세요.')
        sheet=load_workbook(p,read_only=True,data_only=True).active
        rows=[];current=''
        aliases={'경기/충청':'경기/충청','강원동부':'강원동부','강원서부':'강원서부','부산울산':'부산울산'}
        for row in sheet.values:
            v=str(row[0] or '').strip();address=str(row[1] or '').strip() if len(row)>1 else ''
            if v in aliases:current=aliases[v]
            elif current==region and CODE.fullmatch(v.upper()) and address:rows.append((v.upper(),address))
        return rows
    def get_origin(self,region,key):
        if region in self.origins:return self.origins[region]
        self.events.put(('status',region+' 샘플 주소로 좌표 보정 중...'))
        value=calibrate(region,self.sample_rows(region),key)
        if value.status!='사용 가능':
            raise ValueError(f'{region} 샘플의 좌표 오차가 큽니다 (중앙값 {value.median_error_m}m, 최대 {value.maximum_error_m}m). 이 권역은 자동 주소 입력을 중단했습니다.')
        self.origins[region]=value.origin
        self.events.put(('status',f'{region} 보정: {value.sample_count}건, 중앙 오차 {value.median_error_m}m, 최대 {value.maximum_error_m}m'))
        return value.origin
    def fit(self):
        region=self.region.get()
        def job(region,key):
            v=calibrate(region,self.sample_rows(region),key)
            self.events.put(('status',f'{region}: {v.status} / {v.sample_count}건 / 중앙 {v.median_error_m}m / 최대 {v.maximum_error_m}m'))
            if v.status=='사용 가능':self.origins[region]=v.origin
        self.run(job)
    def lookup(self,code,region,key,origin):
        lat,lon=coordinates(code,region,origin)
        return reverse_address(lat,lon,key),lat,lon
    def find(self):
        code=self.code.get().strip().upper()
        if not CODE.fullmatch(code):messagebox.showwarning('번호','전산화번호 8자리를 확인해 주세요.');return
        def job(region,key):
            origin=self.get_origin(region,key)
            address,lat,lon=self.lookup(code,region,key,origin)
            self.events.put(('result',f'{code}  →  {address}\n좌표 {lat:.6f}, {lon:.6f} · 좌표 기반 추정 주소 (현장 확인 권장)'))
        self.run(job)
    def batch(self):
        path=self.input_path.get()
        if not os.path.isfile(path):messagebox.showwarning('엑셀','입력 엑셀 파일을 선택해 주세요.');return
        def job(region,key):
            origin=self.get_origin(region,key)
            book=load_workbook(path);sheet=book.active
            base,_=os.path.splitext(path);out=base+'_독립주소결과.xlsx'
            if os.path.abspath(path)==os.path.abspath(out):raise ValueError('출력 파일 이름이 입력과 같습니다.')
            count=0
            for row in range(1,sheet.max_row+1):
                code=str(sheet.cell(row,1).value or '').strip().upper()
                if not CODE.fullmatch(code):continue
                self.events.put(('status',f'{row}행 {code} 주소 조회 중...'))
                try:
                    address,lat,lon=self.lookup(code,region,key,origin)
                    sheet.cell(row,2).value=address
                    sheet.cell(row,3).value='좌표 기반 추정 주소 — 확인 권장'
                    sheet.cell(row,4).value=lat;sheet.cell(row,5).value=lon
                except Exception as e:
                    sheet.cell(row,2).value='';sheet.cell(row,3).value='조회 실패: '+str(e)[:100]
                count+=1;book.save(out)
            self.events.put(('done',f'{count}건 처리했습니다.\n{out}'))
        self.run(job)
    def start(self):self.root.mainloop()

if __name__=='__main__':App().start()
